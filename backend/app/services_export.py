import io
import re
from PIL import Image as PILImage
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def extraer_nombres(contenido: str):
    firmante = "FIRMANTE AUTORIZADO"
    empleado = "EMPLEADO / DESTINATARIO"
    
    lineas = [l.strip() for l in contenido.split('\n') if l.strip()]
    
    # 1. Buscar emisor tras "Atentamente,"
    for i, l in enumerate(lineas):
        if "Atentamente," in l and i + 1 < len(lineas):
            firmante = lineas[i + 1].strip('.')
            break

    # 2. Buscar empleado en diferentes estructuras formales
    # Patrón A: "Sr.(a) Nombre", "Sr. Nombre", "Sra. Nombre"
    for l in lineas:
        if re.search(r'^(Sr\.\(a\)|Sr\.|Sra\.)', l):
            empleado = re.sub(r'^(Sr\.\(a\)|Sr\.|Sra\.)\s*', '', l).strip('.')
            break
            
    # Patrón B: "empleado(a) Nombre", "empleado Nombre" (Común en Actas Administrativas)
    if empleado == "EMPLEADO / DESTINATARIO":
        match = re.search(r'empleado\(a\)\s+([A-ZÁÉÍÓÚÑa-záéíóúñ\s]+?)(?=\.|\,|$)', contenido)
        if match:
            empleado = match.group(1).strip()

    # Patrón C: "C. Nombre" (Común en Cartas de Recomendación y NDA)
    if empleado == "EMPLEADO / DESTINATARIO":
        match = re.search(r'\bC\.\s+([A-ZÁÉÍÓÚÑa-záéíóúñ\s]+?)(?=\s+laboró|\s+con|\.|\,|$)', contenido)
        if match:
            empleado = match.group(1).strip()

    return firmante, empleado

def exportar_a_pdf(contenido: str, logo_bytes: bytes = None, logo_pos: str = "ninguno") -> bytes:
    buffer = io.BytesIO()
    
    def dibujar_marca_agua(canvas, doc):
        if logo_bytes and logo_pos == "marca_agua":
            try:
                canvas.saveState()
                canvas.setFillAlpha(0.18)
                canvas.setStrokeAlpha(0.18)
                
                pil_img = PILImage.open(io.BytesIO(logo_bytes))
                if pil_img.mode != 'RGBA':
                    pil_img = pil_img.convert('RGBA')
                
                img_reader = ImageReader(pil_img)
                canvas.drawImage(
                    img_reader,
                    x=96,
                    y=186,
                    width=420,
                    height=420,
                    mask='auto',
                    preserveAspectRatio=True
                )
                canvas.restoreState()
            except Exception as e:
                print(f"Error...procesando marca de agua en PDF: {e}")

    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    
    styles = getSampleStyleSheet()
    body_style = ParagraphStyle(
        'LegalBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1e293b'),
        spaceAfter=8
    )

    style_firma_linea = ParagraphStyle(
        'FirmaLinea',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#1e293b')
    )
    
    style_firma_subtexto = ParagraphStyle(
        'FirmaSubtexto',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )
    
    story = []

    # Membrete Superior con Escalado Proporcional
    if logo_bytes and logo_pos in ["arriba_izquierda", "arriba_derecha"]:
        try:
            from reportlab.platypus import Image as RLImage
            pil_img = PILImage.open(io.BytesIO(logo_bytes))
            orig_w, orig_h = pil_img.size
            
            target_width = 140.0
            aspect_ratio = float(orig_h) / float(orig_w)
            target_height = target_width * aspect_ratio

            img_buffer = io.BytesIO(logo_bytes)
            img = RLImage(img_buffer, width=target_width, height=target_height)
            img.hAlign = 'LEFT' if logo_pos == "arriba_izquierda" else 'RIGHT'
            
            story.append(img)
            story.append(Spacer(1, 15))
        except Exception as e:
            print(f"Error...procesando logo membrete: {e}")

    # Párrafos del texto
    for linea in contenido.split('\n'):
        linea_str = linea.strip()
        if linea_str:
            story.append(Paragraph(linea_str, body_style))
        else:
            story.append(Spacer(1, 6))

    # Extraer nombres dinámicos
    nombre_firmante, nombre_empleado = extraer_nombres(contenido)

    # Bloque de Firmas Dinámico
    story.append(Spacer(1, 35))
    tabla_data = [
        [
            Paragraph("___________________________________", style_firma_linea),
            Paragraph("___________________________________", style_firma_linea)
        ],
        [
            Paragraph("Firma del Emisor", style_firma_linea),
            Paragraph("Recibido / Acuse de Enterado", style_firma_linea)
        ],
        [
            Paragraph(nombre_firmante, style_firma_subtexto),
            Paragraph(nombre_empleado, style_firma_subtexto)
        ]
    ]

    tabla_firmas = Table(tabla_data, colWidths=[250, 250])
    tabla_firmas.setStyle(TableStyle([
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ]))
    
    story.append(tabla_firmas)

    if logo_pos == "marca_agua":
        doc.build(story, onFirstPage=dibujar_marca_agua, onLaterPages=dibujar_marca_agua)
    else:
        doc.build(story)

    pdf_data = buffer.getvalue()
    buffer.close()
    return pdf_data


def exportar_a_docx(contenido: str, logo_bytes: bytes = None, logo_pos: str = "ninguno") -> bytes:
    doc = Document()

    if logo_bytes and logo_pos in ["arriba_izquierda", "arriba_derecha"]:
        try:
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT if logo_pos == "arriba_derecha" else WD_ALIGN_PARAGRAPH.LEFT
            logo_buffer = io.BytesIO(logo_bytes)
            p_logo.add_run().add_picture(logo_buffer, width=Inches(1.8))
        except Exception as e:
            print(f"Error procesando logo en Word: {e}")

    for linea in contenido.split('\n'):
        linea_str = linea.strip()
        if linea_str:
            p = doc.add_paragraph(linea_str)
            p.style.font.name = 'Arial'
            p.style.font.size = Pt(10)

    nombre_firmante, nombre_empleado = extraer_nombres(contenido)

    doc.add_paragraph().paragraph_format.space_before = Pt(30)
    p_lineas = doc.add_paragraph("___________________________________           ___________________________________")
    p_lineas.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p_titulos = doc.add_paragraph("Firma del Emisor                                       Recibido / Acuse")
    p_titulos.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_nombres = doc.add_paragraph(f"{nombre_firmante}                                       {nombre_empleado}")
    p_nombres.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_data = buffer.getvalue()
    buffer.close()
    return docx_data